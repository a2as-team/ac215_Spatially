import os
import json
import time
import re
import pandas as pd
import openpyxl
from google import genai
from google.genai import types, errors
from .base import ZoningOrdinanceEmbeddingsBaseProcessor
from utils.gcp_storage import GCPStorage
from .semantic_splitter import SemanticChunker


class BostonZoningOrdinanceEmbeddingsProcessor(ZoningOrdinanceEmbeddingsBaseProcessor):
    def __init__(self):
        super().__init__()
        self.gcp_project = os.environ.get("GCP_PROJECT", "spatially")
        self.gcs_bucket_name = os.environ.get("GCS_BUCKET_NAME")
        self.embedding_model = "text-embedding-004"
        self.embedding_dimension = 768
        self.gcp_location = "us-central1"

        # Initialize LLM client
        self.llm_client = genai.Client(
            vertexai=True,
            project=self.gcp_project,
            location=self.gcp_location
        )

        # Initialize GCS client
        self.storage = GCPStorage(
            gcp_project=self.gcp_project,
            bucket_name=self.gcs_bucket_name
        )

        # Initialize semantic chunker
        self.text_splitter = SemanticChunker(
            embedding_function=self.generate_text_embeddings,
            breakpoint_threshold_type="percentile",
            breakpoint_threshold_amount=95,
            buffer_size=1
        )

        self.logger.info(f"Initialized Boston processor with bucket: {self.gcs_bucket_name}")

    @classmethod
    def city(cls) -> str:
        return "boston"

    def generate_text_embeddings(self, chunks, batch_size=250, max_retries=5, retry_delay=5):
        """Generate embeddings using Vertex AI text-embedding-004"""
        all_embeddings = []

        for i in range(0, len(chunks), batch_size):
            batch = chunks[i:i+batch_size]

            retry_count = 0
            while retry_count <= max_retries:
                try:
                    response = self.llm_client.models.embed_content(
                        model=self.embedding_model,
                        contents=batch,
                        config=types.EmbedContentConfig(
                            output_dimensionality=self.embedding_dimension
                        ),
                    )
                    all_embeddings.extend(
                        [embedding.values for embedding in response.embeddings]
                    )
                    break

                except errors.APIError as e:
                    retry_count += 1
                    if retry_count > max_retries:
                        self.logger.error(
                            f"Failed to generate embeddings after {max_retries} attempts. Last error: {str(e)}"
                        )
                        raise

                    wait_time = retry_delay * (2 ** (retry_count - 1))
                    self.logger.warning(
                        f"API error (code: {e.code}): {e.message}. Retrying in {wait_time} seconds (attempt {retry_count}/{max_retries})..."
                    )
                    time.sleep(wait_time)

        return all_embeddings

    def load_district_codes(self):
        """Load district codes from GCS"""
        try:
            district_codes_path = "zoning_ordinance/district_codes/boston_district_codes.json"
            blob_content = self.storage.download_blob_as_string(district_codes_path)
            return json.loads(blob_content)
        except Exception as e:
            self.logger.warning(f"Could not load district codes: {e}")
            return {}

    def load_district_categories(self):
        """Load district categories from GCS"""
        try:
            categories_path = "zoning_ordinance/district_codes/boston_district_code_categories.json"
            blob_content = self.storage.download_blob_as_string(categories_path)
            return json.loads(blob_content)
        except Exception as e:
            self.logger.warning(f"Could not load district categories: {e}")
            return {}

    def get_code_to_category_mapping(self):
        """Create reverse mapping: district_code -> category_name"""
        categories = self.load_district_categories()
        code_to_category = {}

        for category_name, code_list in categories.items():
            for code in code_list:
                code_to_category[code] = category_name

        return code_to_category

    def extract_district_codes(self, text):
        """Extract district codes from text"""
        if not text:
            return [], []

        district_codes_dict = self.load_district_codes()
        code_to_category = self.get_code_to_category_mapping()

        if not district_codes_dict:
            return [], []

        # Sort codes by length (longest first)
        codes_sorted = sorted(district_codes_dict.keys(), key=len, reverse=True)

        matched_codes = set()
        for code in codes_sorted:
            escaped_code = re.escape(code)

            if len(code) == 1:
                neg_lookbehind = r'(?<!\d-)'
                pattern = neg_lookbehind + r'(?:district|zone|zoning|districts|zones)\s+' + escaped_code + r'\b'
                pattern += r'|' + neg_lookbehind + r'\b' + escaped_code + r'\s+(?:district|zone|zoning|districts|zones)'
                pattern += r'|' + neg_lookbehind + r'(?:,\s*|;\s*)' + escaped_code + r'(?:\s*,|\s*;|\s+)'
            else:
                pattern = r'\b' + escaped_code + r'\b'

            if re.search(pattern, text, re.IGNORECASE):
                matched_codes.add(code)

        # Get categories for matched codes
        matched_categories = set()
        for code in matched_codes:
            category = code_to_category.get(code)
            if category:
                matched_categories.add(category)

        return sorted(list(matched_codes)), sorted(list(matched_categories))

    def extract_structured_data_from_excel(self, excel_path):
        """Extract structured data from Boston Excel files"""
        try:
            workbook = openpyxl.load_workbook(excel_path, data_only=True)
            sheet = workbook.active

            # Get article name from row 4
            article_title = sheet.cell(4, 3).value
            article_subtitle = sheet.cell(4, 4).value

            article_name = ""
            if article_title and article_subtitle:
                article_name = f"{article_title} {article_subtitle}".strip()
            elif article_title:
                article_name = str(article_title).strip()
            elif article_subtitle:
                article_name = str(article_subtitle).strip()

            self.logger.info(f"  Article: {article_name}")

            structured_data = []

            # Process data rows starting from row 3
            for row_num in range(3, sheet.max_row + 1):
                url = sheet.cell(row_num, 1).value
                title = sheet.cell(row_num, 3).value
                subtitle = sheet.cell(row_num, 4).value
                content = sheet.cell(row_num, 5).value

                if not content or not str(content).strip():
                    continue

                section_name = ""
                if title and subtitle:
                    section_name = f"{title} {subtitle}".strip()
                elif title:
                    section_name = str(title).strip()
                elif subtitle:
                    section_name = str(subtitle).strip()

                structured_data.append({
                    "url": str(url) if url else "",
                    "section": section_name,
                    "content": str(content)
                })

            workbook.close()
            return article_name, structured_data

        except Exception as e:
            self.logger.error(f"Error reading Excel file {excel_path}: {e}")
            return "", []

    def process(self, test_mode: bool = False):
        """Process Boston Excel files: chunk and generate embeddings"""
        self.logger.info(f"Starting Boston zoning ordinance embeddings processor (test_mode={test_mode})")

        # List all Excel files in GCS
        input_prefix = self.input_gcs_storage_path()
        self.logger.info(f"Listing files from GCS path: {input_prefix}")

        blobs = self.storage.list_blobs(prefix=input_prefix)
        excel_blobs = [blob for blob in blobs if blob.name.endswith('.xlsx')]

        if test_mode:
            excel_blobs = excel_blobs[:1]

        self.logger.info(f"Found {len(excel_blobs)} Excel files to process")

        for blob in excel_blobs:
            self.logger.info(f"\nProcessing: {blob.name}")

            # Download file to temp location
            local_path = f"/tmp/{os.path.basename(blob.name)}"
            self.storage.download_blob_to_file(blob.name, local_path)

            # Extract structured data
            document_name = os.path.basename(blob.name).rsplit(".", 1)[0]
            article_name, structured_data = self.extract_structured_data_from_excel(local_path)

            self.logger.info(f"Extracted {len(structured_data)} sections")

            all_chunk_data = []

            # Process each section
            for section_data in structured_data:
                content = section_data["content"]
                if not content.strip():
                    continue

                # Chunk the content
                chunks = self.text_splitter.create_documents([content])
                chunk_texts = [doc.page_content for doc in chunks]

                # Create chunk records with metadata
                for chunk_text in chunk_texts:
                    district_codes, district_categories = self.extract_district_codes(chunk_text)

                    all_chunk_data.append({
                        "chunk": chunk_text,
                        "document": document_name,
                        "city": "boston",
                        "article": article_name,
                        "section": section_data["section"],
                        "url": section_data["url"],
                        "district_code": district_codes,
                        "district_category": district_categories
                    })

            if not all_chunk_data:
                self.logger.warning(f"No chunks created for {document_name}")
                continue

            # Generate embeddings
            self.logger.info(f"Generating embeddings for {len(all_chunk_data)} chunks")
            data_df = pd.DataFrame(all_chunk_data)

            chunks = data_df["chunk"].values.tolist()
            embeddings = self.generate_text_embeddings(chunks, batch_size=15)
            data_df["embedding"] = embeddings

            # Save to GCS as JSONL
            output_path = f"{self.output_gcs_storage_path()}/embeddings-{document_name}.jsonl"
            jsonl_content = data_df.to_json(orient='records', lines=True)

            self.storage.upload_string_to_blob(jsonl_content, output_path)
            self.logger.info(f"Uploaded embeddings to: {output_path}")

            # Cleanup
            os.remove(local_path)

            # Rate limiting
            time.sleep(2)

        self.logger.info("Boston embeddings processing complete!")
